from io import BytesIO
import os
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from core import storage
from backtesting.s18 import service
from backtesting.s18.book import Book, Position
from backtesting.s18.config import S18Config
from backtesting.s18.replay import validate_dataset
from backtesting.s18.dataset import get_dataset
from test_s18_book import market


@pytest.fixture
def daily_setup(monkeypatch, tmp_path):
    from backtesting.s18 import data

    monkeypatch.setenv("PORTFOLIO_DB_PATH", str(tmp_path / "test.sqlite3"))
    monkeypatch.setattr(service, "REPORTS_ROOT", tmp_path / "reports")
    m = market(rows=150)
    m.provenance = {
        "session_calendar": [str(d) for d in m.dates],
        "price_scales": {},
        "history_fingerprints": {},
    }
    monkeypatch.setattr(
        service,
        "get_dataset",
        lambda: SimpleNamespace(id="fixture", status=lambda: {"dataset_id": "fixture"}),
    )
    monkeypatch.setattr(
        service,
        "_certificate",
        lambda kit, required: {
            "status": "passed",
            "source_hash": "code",
            "data_hash": "data",
            "results": [],
        },
    )
    monkeypatch.setattr(data, "load_market", lambda *a, **kw: m)
    return m


def test_daily_idempotence_catchup_and_shadow_costs(daily_setup):
    m = daily_setup
    first = service.run_daily(as_of=str(m.dates[0]), combo="P5", metal_mode="none")
    assert first["holdings"] == []
    assert len(first["orders"]) == 10
    caught_up = service.run_daily(as_of=str(m.dates[5]), combo="P5", metal_mode="none")
    repeated = service.run_daily(as_of=str(m.dates[5]), combo="P5", metal_mode="none")
    assert repeated["fills"] == caught_up["fills"]
    assert repeated["equity_curve"] == caught_up["equity_curve"]
    assert len(repeated["holdings"]) == 10
    charges = service.run_daily(
        as_of=str(m.dates[5]),
        combo="P5",
        metal_mode="none",
        shadow_charges={"A-1": {"stt": 5, "stamp_duty": 1, "gst": 2}},
    )
    assert charges["fills"][0]["actual_charges"] == 8
    assert charges["metrics"] == repeated["metrics"]
    assert charges["holdings"] == repeated["holdings"]
    with pytest.raises(ValueError, match="Unknown model fill"):
        service.run_daily(
            as_of=str(m.dates[5]),
            combo="P5",
            metal_mode="none",
            shadow_charges={"not-a-fill": {"stt": 1}},
        )
    with pytest.raises(ValueError, match="backwards"):
        service.run_daily(as_of=str(m.dates[3]), combo="P5", metal_mode="none")


def test_daily_dryrun_never_persists_book(daily_setup):
    result = service.run_daily(as_of="2025-01-01", persist=False)
    assert not result["state_persisted"]
    assert storage.get_document("s18", "book:default:P15:both_priority") is None


def test_native_pending_session_never_creates_paper_positions(daily_setup, monkeypatch):
    from backtesting.s18 import live_data

    monkeypatch.setattr(
        live_data,
        "collect_market",
        lambda *a, **k: SimpleNamespace(
            market=daily_setup,
            readiness={
                "status": "waiting_for_first_snapshot",
                "message": "Waiting for a dated snapshot",
                "latest_price_date": "2026-10-01",
            },
        ),
    )
    result = service.run_daily(as_of="2026-10-04", capital=500000)
    assert not result["state_persisted"]
    assert result["orders"] == []
    assert result["metrics"]["planned_capital"] == 500000
    assert storage.get_document("s18", "book:default:P15:both_priority") is None


def test_native_ready_book_starts_from_cash_and_fills_next_session(
    daily_setup, monkeypatch
):
    from backtesting.s18 import live_data

    m = market(rows=45, start="2026-09-01")
    m.provenance = {
        "session_calendar": [str(d) for d in m.dates],
        "calendar_complete_through": "2026-11-02",
        "as_of": "2026-09-03",
        "price_scales": {},
        "history_fingerprints": {},
    }
    calls = []

    def collect(*args, **kwargs):
        calls.append(kwargs)
        return SimpleNamespace(market=m, readiness={"status": "ready"})

    monkeypatch.setattr(live_data, "collect_market", collect)
    initial = service.run_daily(
        as_of="2026-09-03",
        capital=500000,
        book_id="native-forward",
        metal_mode="none",
    )
    assert initial["metrics"]["final_value"] == 500000
    assert initial["holdings"] == []
    assert initial["equity_curve"][0]["date"] == "2026-09-03"
    assert len(initial["orders"]) == 30
    assert all(order["paper_only"] for order in initial["orders"])
    m.provenance["as_of"] = "2026-09-04"
    continued = service.run_daily(
        as_of="2026-09-04",
        capital=500000,
        book_id="native-forward",
        metal_mode="none",
    )
    assert calls[-1]["inception"] == "2026-09-03"
    assert len(continued["holdings"]) == 30
    assert len(continued["fills"]) == 30
    assert continued["metrics"]["final_value"] == pytest.approx(499900)
    assert continued["portfolio_state"]["inception_date"] == "2026-09-03"


def test_certificates_for_two_runtimes_do_not_overwrite_each_other(
    monkeypatch, tmp_path
):
    monkeypatch.setenv("PORTFOLIO_DB_PATH", str(tmp_path / "cert.sqlite3"))
    monkeypatch.setattr(service, "source_fingerprint", lambda: "scheduler-runtime")
    monkeypatch.setattr(service, "input_fingerprint", lambda p: "data")
    monkeypatch.setattr(service, "reference_fingerprint", lambda p: "reference")
    valid = {
        "status": "passed",
        "source_hash": "scheduler-runtime",
        "data_hash": "data",
        "reference_hash": "reference",
    }
    storage.set_document("s18", "golden_replay:scheduler-runtime", valid)
    storage.set_document(
        "s18", "golden_replay", {"status": "passed", "source_hash": "ui-runtime"}
    )
    assert service._certificate(tmp_path, required=True) == valid


def test_daily_refuses_stale_live_data_and_mutated_capital(daily_setup, monkeypatch):
    from backtesting.s18 import live_data

    def unavailable(**kwargs):
        raise ValueError("NSE data unavailable; refusing stale prices")

    monkeypatch.setattr(live_data, "collect_market", unavailable)
    with pytest.raises(ValueError, match="refusing stale"):
        service.run_daily(as_of="2026-09-01")
    service.run_daily(as_of="2025-01-01")
    with pytest.raises(ValueError, match="capital cannot change"):
        service.run_daily(as_of="2025-01-02", capital=200000)


def test_resume_rejects_revised_inputs(daily_setup):
    m = daily_setup
    m.provenance["history_fingerprints"] = {"2025-01-01": "original"}
    service.run_daily(as_of="2025-01-01")
    m.provenance["history_fingerprints"]["2025-01-01"] = "revised"
    with pytest.raises(ValueError, match="forward data was revised"):
        service.run_daily(as_of="2025-01-02")


def test_atomic_book_compare_and_swap(monkeypatch, tmp_path):
    monkeypatch.setenv("PORTFOLIO_DB_PATH", str(tmp_path / "test.sqlite3"))
    service._save_book("test", None, {"v": 1})
    with pytest.raises(ValueError, match="changed during this run"):
        service._save_book("test", None, {"v": 2})
    assert storage.get_document("s18", "test") == {"v": 1}


def test_new_ipo_remaps_existing_metals_by_identity():
    book = Book(S18Config("P20", "both"), "A")
    book.positions = [Position(2, 1, 100, 100, 0, 0, "S1", "S1")]
    book.queued_buys = [{"gid": 3, "weight": 2, "sleeve": "S1"}]
    service._remap_books(
        [book],
        ["stock0", "stock1", "gold", "silver"],
        ["stock0", "stock1", "new_ipo", "gold", "silver"],
    )
    assert book.positions[0].gid == 3
    assert book.queued_buys[0]["gid"] == 4
    with pytest.raises(ValueError, match="dropped an existing"):
        service._remap_books([book], ["gold", "silver"], ["gold"])


def test_current_quotes_convert_without_mutating_model_book():
    m = market(rows=3)
    m.provenance["raw_to_book"] = {identity: 2.0 for identity in m.identities}
    book = Book(S18Config("P5", "none"), "A")
    book.last_row, book.last_date = 1, str(m.dates[1])
    book.equity = [{"nav": 101}]
    p = Position(
        0, 1, 80, 120, 0, 0, "S1", "S1", armed=True, queued_sell_reason="floor"
    )
    book.positions = [p]
    holding = service._display_holdings([book], m)[0]
    order = service._orders([book], m, str(m.dates[2]))[0]
    assert holding["qty"] == 2 and holding["model_qty"] == 1
    assert holding["close"] == 50 and holding["peak"] == 60
    assert holding["stop"] == 48
    assert holding["qty"] * holding["close"] == holding["value"]
    assert order["qty"] == 2 and order["paper_only"]
    assert p.qty == 1 and p.peak == 120 and p.basis == 80
    del m.provenance["raw_to_book"][m.identities[0]]
    with pytest.raises(ValueError, match="Missing current quote conversion"):
        service._orders([book], m, str(m.dates[2]))


def test_rebase_trade_identity_survives_ticker_reuse():
    m = market(rows=3)
    book = Book(S18Config("P5", "none"), "A")
    book.trades = [
        {"symbol": "REUSED", "entry_px": 100.0, "exit_px": 110.0},
        {"symbol": "REUSED", "entry_px": 200.0, "exit_px": 220.0},
    ]
    book.trade_identities = ["first-company", "second-company"]
    service._rescale_books(
        [book],
        {},
        {"first-company": 0.5, "second-company": 0.25},
        ["first-company", "second-company"],
        m,
    )
    assert [t["entry_px"] for t in book.trades] == [50.0, 50.0]
    assert [t["exit_px"] for t in book.trades] == [55.0, 55.0]


def test_certificate_invalidated_when_engine_changes(monkeypatch, tmp_path):
    monkeypatch.setenv("PORTFOLIO_DB_PATH", str(tmp_path / "test.sqlite3"))
    monkeypatch.setattr(service, "source_fingerprint", lambda: "new")
    storage.set_document(
        "s18",
        "golden_replay",
        {
            "status": "passed",
            "source_hash": "old",
            "data_hash": "data",
        },
    )
    with pytest.raises(ValueError, match="golden replay"):
        service._certificate(tmp_path, required=True)


def test_backtest_persists_dossier(daily_setup):
    m = daily_setup
    result = service.run_backtest(
        start=str(m.dates[1]),
        end=str(m.dates[-1]),
        combo="P5",
        metal_mode="none",
    )
    artifact = storage.get_artifact(result["artifact_group_id"], "s18_dossier.xlsx")
    assert (
        Path(result["results_dir"]) / "s18_dossier.xlsx"
    ).read_bytes() == artifact.payload
    workbook = load_workbook(BytesIO(artifact.payload), read_only=True)
    assert "Tax_Ledger" in workbook.sheetnames
    assert "Replay_Validation" in workbook.sheetnames
    assert workbook["Equity_Curve"].max_row == len(m.dates) + 1
    assert workbook["Trades"].max_row == len(result["trades"]) + 1
    assert workbook["Positions"].max_row == len(result["holdings"]) + 1


def test_rebase_preserves_value_and_basis(daily_setup):
    m = daily_setup
    service.run_daily(as_of=str(m.dates[0]), combo="P5", metal_mode="none")
    before = service.run_daily(as_of=str(m.dates[1]), combo="P5", metal_mode="none")
    m.provenance["price_scales"] = {identity: 0.5 for identity in m.identities}
    for values in (m.open, m.high, m.low, m.close):
        values *= 0.5
    after = service.run_daily(as_of=str(m.dates[2]), combo="P5", metal_mode="none")
    for a, b in zip(after["holdings"], before["holdings"]):
        assert a["qty"] == pytest.approx(b["qty"] * 2)
        assert a["basis"] == b["basis"]
        assert a["value"] == pytest.approx(b["value"])
        assert a["peak"] == pytest.approx(b["peak"] / 2)


def test_all_s18_golden_references():
    db = os.environ.get("S18_DATASET_DB")
    if not db:
        pytest.skip("Set S18_DATASET_DB to replay the installed historical dataset")
    certificate, artifacts = validate_dataset(get_dataset(db_path=Path(db)))
    assert certificate["status"] == "passed"
    assert certificate["tranches_checked"] == 30
    assert len(artifacts) == 30
    assert max(r["sharpe_delta"] for r in certificate["results"]) < 0.001
    assert max(r["max_nav_error"] for r in certificate["results"]) < 1e-7
