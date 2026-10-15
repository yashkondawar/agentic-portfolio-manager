from copy import deepcopy
from io import BytesIO

import numpy as np
from openpyxl import load_workbook
import pytest

from backtesting.s18 import historical, portfolio, service
from backtesting.s18.book import Book, Position
from backtesting.s18.config import S18Config
from core import storage
from test_s18_book import step
from test_s18_service import daily_setup  # noqa: F401
from test_s18_whole_shares import live_market, native_setup  # noqa: F401


def bonus_book(*, shares=112.5, action="bonus", factor=2 / 3, start="2025-03-31"):
    m = live_market(rows=5, start=start)
    m.provenance["raw_to_book_events"] = [{
        "date": str(m.dates[1]), "identity": m.identities[0],
        "kind": action, "factor": factor,
    }]
    book = Book(S18Config("P5", "none"), "A", cash=300, whole_shares=True)
    book.last_row, book.last_date = 0, str(m.dates[0])
    book.positions = [Position(0, shares, shares * 80, 100, 0, 0, "S1", "S1")]
    book.equity = [{"nav": 300 + shares * 100}]
    return book, m


def test_bonus_fraction_is_cash_not_a_fractional_exchange_fill():
    book, m = bonus_book()
    step(book, m, 1, decisions=False)
    assert book.positions[0].qty == 112
    assert book.positions[0].basis == pytest.approx(8960)
    assert book.cash == pytest.approx(350)
    assert book.equity[-1]["nav"] == pytest.approx(11550)
    assert not book.pending_cash_in_lieu and not book.fills and not book.trades
    settlement, = book.cash_in_lieu
    assert settlement["share_qty"] == pytest.approx(0.5)
    assert settlement["amount"] == pytest.approx(50)
    assert settlement["cost_basis"] == pytest.approx(40)
    assert settlement["gain"] == pytest.approx(10)
    assert settlement["model_cost"] == book.model_cost == 0
    assert settlement["modeled"]
    assert book.realised_st == pytest.approx(10)
    assert service._display_holdings([book], m)[0]["qty"] == 112


def test_settlement_gains_follow_tax_rollover_not_previous_year():
    book, m = bonus_book()
    book.realised_st = 100
    step(book, m, 1, decisions=False)
    assert book.tax_paid == 20
    assert book.cash == pytest.approx(330)
    assert book.realised_st == pytest.approx(10)
    assert book.tax_ledger[0]["realised_st"] == 100


def test_fraction_can_consume_the_entire_position_after_consolidation():
    book, m = bonus_book(shares=0.5, action="split", factor=2)
    step(book, m, 1, decisions=False)
    assert not book.positions
    assert book.cash == pytest.approx(350)
    assert book.cash_in_lieu[0]["cost_basis"] == pytest.approx(40)
    assert not book.fills


def test_missing_open_keeps_receivable_and_restart_settles_it_once():
    book, m = bonus_book()
    m.tradable[1:3, 0] = False
    step(book, m, 1, decisions=False)
    assert book.cash == 300
    assert book.equity[-1]["nav"] == pytest.approx(11550)
    assert book.equity[-1]["pending_cash_in_lieu"] == pytest.approx(50)
    assert not book.cash_in_lieu
    assert book.entitlement_rows(m)[0]["entitlement_shares"] == pytest.approx(0.5)
    book = Book.from_dict(book.to_dict())
    step(book, m, 2, decisions=False)
    assert len(book.pending_cash_in_lieu) == 1
    m.open[3, 0] = 104
    step(book, m, 3, decisions=False)
    assert book.cash == pytest.approx(352)
    assert book.cash_in_lieu[0]["ex_date"] == str(m.dates[1])
    assert book.cash_in_lieu[0]["date"] == str(m.dates[3])
    step(book, m, 4, decisions=False)
    assert book.cash == pytest.approx(352)
    assert len(book.cash_in_lieu) == 1


def test_pending_receivable_is_not_credited_by_close_price_order_preview():
    book, m = bonus_book()
    m.tradable[1, 0] = False
    step(book, m, 1, decisions=False)
    preview = deepcopy(book)
    m.tradable[1, 0] = True
    preview._open(m, 1, at_close=True, execution_date=str(m.dates[2]))
    assert preview.cash == 300
    assert preview.pending_cash_in_lieu == book.pending_cash_in_lieu
    assert not preview.cash_in_lieu


def test_pending_units_rebase_and_remap_without_changing_basis():
    book, m = bonus_book()
    m.tradable[1, 0] = False
    step(book, m, 1, decisions=False)
    service._rescale_books(
        [book], {}, {m.identities[0]: 0.5}, list(m.identities), m
    )
    assert book.pending_cash_in_lieu[0].qty == pytest.approx(1)
    assert book.pending_cash_in_lieu[0].basis == pytest.approx(40)
    assert book.pending_cash_in_lieu[0].original_share_qty == pytest.approx(0.5)
    reordered = list(m.identities)
    reordered[0], reordered[1] = reordered[1], reordered[0]
    service._remap_books([book], list(m.identities), reordered)
    assert book.pending_cash_in_lieu[0].gid == book.positions[0].gid == 1


def test_second_split_before_first_quote_preserves_entitlement_value():
    book, m = bonus_book()
    m.execution_factors = np.ones(m.close.shape)
    m.execution_factors[0, 0] = 2 / 3
    m.execution_factors[2:, 0] = 2
    m.provenance["raw_to_book_events"].append({
        "date": str(m.dates[2]), "identity": m.identities[0],
        "kind": "split", "factor": 0.5,
    })
    m.tradable[1, 0] = False
    step(book, m, 1, decisions=False)
    step(book, m, 2, decisions=False)
    assert service._display_holdings([book], m)[0]["qty"] == 224
    assert book.cash_in_lieu[0]["share_qty"] == pytest.approx(1)
    assert book.cash_in_lieu[0]["share_price"] == pytest.approx(50)
    assert book.cash_in_lieu[0]["amount"] == pytest.approx(50)


def test_pending_settlement_rebase_matches_one_shot_live_catchup():
    incremental, prior = bonus_book()
    prior.tradable[1, 0] = False
    full_book = deepcopy(incremental)
    current = deepcopy(prior)
    for prices in (current.open, current.high, current.low, current.close):
        prices[:, 0] /= 2
    current.provenance["raw_to_book_events"].append({
        "date": str(current.dates[2]), "identity": current.identities[0],
        "kind": "split", "factor": 0.5,
    })
    full_book.positions[0].qty *= 2
    full_book.positions[0].peak /= 2
    step(full_book, current, 1, decisions=False)
    step(full_book, current, 2, decisions=False)
    step(incremental, prior, 1, decisions=False)
    incremental = Book.from_dict(incremental.to_dict())
    service._rescale_books(
        [incremental], {}, {current.identities[0]: 0.5}, list(current.identities), current
    )
    step(incremental, current, 2, decisions=False)
    assert incremental.to_dict() == full_book.to_dict()


def test_unexplained_fraction_still_fails_instead_of_being_paid():
    book, m = bonus_book()
    m.provenance["raw_to_book_events"] = []
    with pytest.raises(ValueError, match="Fractional share entitlement"):
        step(book, m, 1, decisions=False)
    assert book.cash == 300 and not book.cash_in_lieu


def test_demerger_of_pending_entitlement_is_not_silently_settled():
    book, m = bonus_book(shares=0.5, action="split", factor=2)
    m.tradable[1, 0] = False
    step(book, m, 1, decisions=False)
    m.provenance["raw_to_book_events"].append({
        "date": str(m.dates[2]), "identity": m.identities[0],
        "kind": "demerger", "factor": 1,
    })
    with pytest.raises(ValueError, match="demerger requires holdings reconciliation"):
        step(book, m, 2, decisions=False)
    assert not book.cash_in_lieu


@pytest.mark.parametrize("shares", [112.0, 112.0 + 5e-8])
def test_integer_or_rounding_noise_creates_no_cash_receivable(shares):
    book, m = bonus_book(shares=shares)
    step(book, m, 1, decisions=False)
    assert not book.pending_cash_in_lieu and not book.cash_in_lieu
    assert book.cash == 300


def test_live_catchup_persists_pending_nav_and_matches_backtest(request, monkeypatch):
    m = request.getfixturevalue("native_setup")
    m.provenance["raw_to_book_events"] = [{
        "date": "2026-09-07", "identity": m.identities[0], "kind": "bonus", "factor": 2 / 3,
    }]
    m.tradable[4, 0] = False
    params = dict(book_id="cash-lieu", combo="P5", metal_mode="none", capital=5000)
    service.run_daily(**params, as_of="2026-09-03")
    m.provenance["as_of"] = "2026-09-07"
    pending = service.run_daily(**params, as_of="2026-09-07")
    assert pending["pending_cash_in_lieu"]
    snap = portfolio.ledger_snapshot(book_id="cash-lieu", combo="P5", metal_mode="none")
    assert snap["book"]["pending_entitlement_value"] == pytest.approx(100)
    assert snap["book"]["equity"] == pytest.approx(
        snap["book"]["cash"] + snap["book"]["deployed"] + 100
    )
    m.provenance["as_of"] = "2026-09-08"
    settled = service.run_daily(**params, as_of="2026-09-08")
    assert not settled["pending_cash_in_lieu"]
    assert len(settled["cash_in_lieu"]) == 2
    again = service.run_daily(**params, as_of="2026-09-08")
    assert again["portfolio_state"]["books"] == settled["portfolio_state"]["books"]
    monkeypatch.setattr(historical, "load_execution_market", lambda **kw: m)
    saved = storage.get_document("s18", "book:cash-lieu:P5:none")
    result = service.run_backtest(
        combo="P5", metal_mode="none", capital=5000,
        start="2026-09-03", end="2026-09-08",
    )
    assert result["blocked"] is None
    assert result["portfolio_state"]["books"] == saved["books"]
    assert storage.get_document("s18", "book:cash-lieu:P5:none") == saved
    assert result["comparison"][1]["modeled_settlements"] == 2
    assert all(isinstance(f["qty"], int) for f in result["fills"])
    artifact = storage.get_artifact(result["artifact_group_id"], "s18_dossier.xlsx")
    workbook = load_workbook(BytesIO(artifact.payload), read_only=True)
    assert workbook["Modeled_Cash_In_Lieu"].max_row == 3
    assert "Pending_Entitlements" in workbook.sheetnames
