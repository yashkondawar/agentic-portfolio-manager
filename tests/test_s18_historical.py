from copy import deepcopy
from io import BytesIO
import json
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest
from streamlit.testing.v1 import AppTest

from backtesting.s18 import historical, service
from backtesting.s18.config import S18Config
from backtesting.s18.data import PRICE_FIELDS, RawMarket
from backtesting.s18.sizing import ExecutionDataError, quote_factor
from core import storage
from strategies.s18 import _result
from test_s18_service import daily_setup  # noqa: F401
from test_s18_whole_shares import native_setup  # noqa: F401


def test_historical_execution_matches_live_catchup_exactly(request):
    m = request.getfixturevalue("native_setup")
    params = dict(book_id="parity", combo="P5", metal_mode="none", capital=500000)
    service.run_daily(**params, as_of="2026-09-03")
    m.provenance["as_of"] = "2026-09-10"
    daily = service.run_daily(**params, as_of="2026-09-10")
    simulated = historical.simulate_live(
        m, S18Config("P5", "none"), start="2026-09-03", end="2026-09-10",
        capital_per_tranche=250000,
    )
    assert [b.to_dict() for b in simulated] == daily["portfolio_state"]["books"]


def test_historical_startup_has_no_forced_review_or_early_scout(request):
    m = request.getfixturevalue("native_setup")
    m.risk_off[:] = True
    books = historical.simulate_live(
        m, S18Config("P5", "none"), start="2026-09-03", end="2026-09-10",
        capital_per_tranche=250000,
    )
    assert all(not b.fills and not b.positions for b in books)
    assert books[0].equity[0]["date"] == "2026-09-03"


def test_blocked_backtest_never_labels_prefix_as_full_result(request, monkeypatch):
    m = request.getfixturevalue("native_setup")
    m.provenance["raw_to_book_events"] = [{
        "date": "2026-09-07", "identity": m.identities[0], "kind": "demerger", "factor": 1.0,
    }]
    monkeypatch.setattr(historical, "load_execution_market", lambda **kw: m)
    output = service.run_backtest(
        combo="P5", metal_mode="none", start="2026-09-03", end="2026-09-10",
        capital=500000,
    )
    assert output["blocked"]["date"] == "2026-09-07"
    assert output["blocked"]["completed_through"] == "2026-09-04"
    assert output["blocked"]["symbol"] == m.symbols[0]
    assert output["metrics"] == {}
    assert output["partial_metrics"]["end_date"] == "2026-09-04"
    assert {r["end_date"] for r in output["comparison"]} == {"2026-09-04"}
    assert all(f["date"] < "2026-09-07" for f in output["fills"])
    assert all(isinstance(f["qty"], int) for f in output["fills"])
    assert "INCOMPLETE" in output["report"]
    result = _result("s18_backtest", output)
    assert not result.ok
    assert "2026-09-07" in result.error
    assert storage.get_document("s18", "book:default:P5:none") is None
    artifact = storage.get_artifact(output["artifact_group_id"], "s18_dossier.xlsx")
    workbook = load_workbook(BytesIO(artifact.payload), read_only=True)
    assert "Sizing_Comparison" in workbook.sheetnames
    assert any("INCOMPLETE" in str(row[1].value) for row in workbook["Summary"].rows)
    app = AppTest.from_string(
        "import json\n"
        "from core.strategy import StrategyResult\n"
        "from ui.components import render_result\n"
        f"render_result(StrategyResult(**json.loads({json.dumps(result.to_dict())!r})))\n"
    ).run(timeout=60)
    assert not app.exception
    assert any("Backtest stopped" in error.value for error in app.error)
    assert any("Partial-period comparison only" in caption.value for caption in app.caption)


def test_matrix_quote_units_are_verified_on_the_actual_session(request):
    m = request.getfixturevalue("native_setup")
    m.execution_factors = np.full(m.close.shape, 0.5)
    assert quote_factor(m, m.identities[0], 1) == 0.5
    m.execution_factors[2, 0] = np.nan
    with pytest.raises(ExecutionDataError, match="2026-09-03"):
        quote_factor(m, m.identities[0], 2)


def raw_fixture(*, metals=False):
    count = 4 if metals else 2
    dates = pd.bdate_range("2024-01-01", periods=300).to_numpy(dtype="datetime64[D]")
    prices = {f: np.full((len(dates), count), 100.0) for f in PRICE_FIELDS}
    identities = tuple(f"identity-{i}" for i in range(count))
    raw = RawMarket(
        dates, deepcopy(prices), np.ones(prices["close"].shape, dtype=bool),
        np.full(len(dates), 100.0), np.arange(len(dates)), identities, identities,
        tuple((f"isin-{i}",) for i in range(count)), 2,
        {"session_calendar": [str(d) for d in dates], "calendar_complete_through": str(dates[-1])},
    )
    return raw, prices


def test_quote_normalization_does_not_invent_fractional_entitlements(monkeypatch):
    raw, prices = raw_fixture()
    for field in PRICE_FIELDS:
        raw.prices[field] *= 0.123456789
        raw.prices[field][1::2] *= 1 + 2e-8
        raw.prices[field][5] = np.nan
        prices[field][5] = np.nan
    monkeypatch.setattr(historical, "load_raw_market", lambda **kw: raw)
    monkeypatch.setattr(historical, "_raw_quotes", lambda *a: (prices, [], {}))
    m = historical.load_execution_market(source=SimpleNamespace(), metals=False)
    assert np.ptp(m.execution_factors[:, 0]) == 0
    assert not m.tradable[5, 0]
    assert m.execution_factors[5, 0] == m.execution_factors[4, 0]
    assert m.open[1, 0] / m.execution_factors[1, 0] == pytest.approx(100)


def test_unverified_price_unit_changes_are_flagged_not_assumed_splits(monkeypatch):
    raw, prices = raw_fixture()
    for field in PRICE_FIELDS:
        raw.prices[field][20:, 0] *= 2
    monkeypatch.setattr(historical, "load_raw_market", lambda **kw: raw)
    monkeypatch.setattr(historical, "_raw_quotes", lambda *a: (prices, [], {}))
    m = historical.load_execution_market(source=SimpleNamespace(), metals=False)
    assert m.provenance["raw_to_book_events"] == [{
        "date": str(raw.dates[20]), "identity": raw.identities[0],
        "kind": "unverified", "factor": 0.5,
    }]


@pytest.mark.parametrize("missing_source", ["exchange", "adjusted"])
def test_stocks_cannot_trade_without_complete_verified_ohlc(monkeypatch, missing_source):
    raw, prices = raw_fixture()
    source = prices if missing_source == "exchange" else raw.prices
    source["open"][20, 0] = np.nan
    monkeypatch.setattr(historical, "load_raw_market", lambda **kw: raw)
    monkeypatch.setattr(historical, "_raw_quotes", lambda *a: (prices, [], {}))
    m = historical.load_execution_market(source=SimpleNamespace(), metals=False)
    assert not m.tradable[20, 0]
    assert not raw.universe[20, 0]
    assert np.isnan(raw.prices["close"][20, 0])
    assert m.execution_factors[20, 0] == 1


def test_split_on_nontrading_day_updates_units_without_inventing_a_quote(monkeypatch):
    raw, prices = raw_fixture()
    for field in PRICE_FIELDS:
        raw.prices[field][20:, 0] *= 2
        raw.prices[field][20, 0] = np.nan
        prices[field][20, 0] = np.nan
    actions = [("isin-0", str(raw.dates[20]), "split", 0.5, "Split 2:1")]
    monkeypatch.setattr(historical, "load_raw_market", lambda **kw: raw)
    monkeypatch.setattr(
        historical, "_raw_quotes", lambda *a: (prices, actions, {"isin-0": [0]})
    )
    m = historical.load_execution_market(source=SimpleNamespace(), metals=False)
    assert not m.tradable[20, 0]
    assert m.execution_factors[19, 0] == 1
    assert m.execution_factors[20, 0] == m.execution_factors[21, 0] == 2
    assert m.provenance["raw_to_book_events"][0]["kind"] == "split_bonus"


def test_proxy_etf_period_is_removed_from_prices_and_signal_warmup(monkeypatch):
    raw, prices = raw_fixture(metals=True)
    for field in PRICE_FIELDS:
        prices[field][:50, 3] = np.nan
    monkeypatch.setattr(historical, "load_raw_market", lambda **kw: raw)
    monkeypatch.setattr(historical, "_raw_quotes", lambda *a: (prices, [], {}))
    m = historical.load_execution_market(source=SimpleNamespace(), metals=True)
    assert np.isnan(raw.prices["close"][:50, 3]).all()
    assert not m.tradable[:50, 3].any()
    assert not m.x63[:299, 3].any()
    assert m.provenance["metal_availability"][1]["first_verified_quote"] == str(raw.dates[50])


@pytest.mark.parametrize("missing_quote", [False, True])
def test_simultaneous_split_and_bonus_are_combined_once(monkeypatch, missing_quote):
    raw, prices = raw_fixture()
    for field in PRICE_FIELDS:
        raw.prices[field][20:, 0] *= 10
        if missing_quote:
            raw.prices[field][20, 0] = prices[field][20, 0] = np.nan
    actions = [
        ("isin-0", str(raw.dates[20]), "split", 0.5, "Split 2:1"),
        ("isin-0", str(raw.dates[20]), "bonus", 0.2, "Bonus 4:1"),
    ]
    monkeypatch.setattr(historical, "load_raw_market", lambda **kw: raw)
    monkeypatch.setattr(historical, "_raw_quotes", lambda *a: (prices, actions, {"isin-0": [0]}))
    m = historical.load_execution_market(source=SimpleNamespace(), metals=False)
    assert m.execution_factors[20, 0] == m.execution_factors[21, 0] == 10
    assert m.provenance["raw_to_book_events"] == [{
        "date": str(raw.dates[20]), "identity": raw.identities[0],
        "kind": "split_bonus", "factor": 0.1,
    }]


@pytest.mark.parametrize("entries", [
    [("split", 0.5, ""), ("split", 0.5, "")],
    [("split_bonus", 0.1, ""), ("bonus", 0.2, "")],
    [("bonus", float("nan"), "")],
])
def test_duplicate_or_overlapping_actions_are_not_assumed_valid(entries):
    assert historical._combined_share_factor(entries) is None


@pytest.mark.parametrize("kind,factor", [
    ("split+bonus", 0.25), ("bonus+dividend", 0.5), ("split+dividend", 0.1),
])
def test_historical_kind_aliases_preserve_verified_share_factors(kind, factor):
    assert historical._combined_share_factor([(kind, factor, "")]) == factor
